#streamlit-> a python library to build UI
import streamlit as st
#pdfplumber-> exports text from PDF reliably
import pdfplumber
import requests
#Document() lets u read paras and text
from docx import Document
from dotenv import load_dotenv

load_dotenv()

# ===== KEYWORD EXTRACTION =====
def extract_keyword(jd_text):
    prompt = f"""
You are a job analyzer.

Extract ONLY important technical skills from the job description.

Rules:
- Return ONLY comma-separated keywords
- No sentences
- No explanation
- Only skills/tools/technologies

Example output:
Python, Machine Learning, SQL, TensorFlow

Job Description:
{jd_text}
"""
    response = requests.post(
        "http://localhost:11434/api/generate",
        json={
            "model": "mistral",
            "prompt": prompt,
            "stream": False
        }
    )
    data = response.json()
    return data.get("response", "")


# ===== RESUME OPTIMIZATION =====
def call_local_llm(resume_text, jd_text):
    prompt = f"""
You are an expert ATS resume writer.

TASK:
Rewrite and optimize the resume to perfectly match the job description.

STRICT RULES:
- Output ONLY the final resume
- No explanations, no notes, no "Dear user"
- Do NOT mention AI
- Use clear section headings
- Use bullet points
- Keep it concise and professional

STRUCTURE:
1. Summary
2. Skills
3. Experience
4. Projects
5. Education

INPUT RESUME:
{resume_text}

JOB DESCRIPTION:
{jd_text}
"""
    response = requests.post(
        "http://localhost:11434/api/generate",
        json={
            "model": "mistral",
            "prompt": prompt,
            "stream": False
        }
    )
    data = response.json()
    return data.get("response", "")


# ===== UI =====
st.set_page_config(page_title="AI Resume Optimizer", layout="wide")
st.title("AI Resume Optimizer")

st.write("Upload your resume and paste the job description to optimize it.")

col1, col2 = st.columns(2)

with col1:
    uploaded_file = st.file_uploader("Upload Resume (PDF/DOCX)", type=["pdf", "docx"])

with col2:
    role = st.text_input("Target Role (optional)")

jd_text = st.text_area("Paste Job Description", height=200)


# ===== MAIN LOGIC =====
if st.button("Optimize Resume"):

    if uploaded_file is None or jd_text.strip() == "":
        st.warning("Please upload resume and enter job description")

    else:
        st.write("Parsing resume...")

        # ===== PARSE RESUME =====
        if uploaded_file.name.endswith(".pdf"):
            with pdfplumber.open(uploaded_file) as pdf:
                resume_text = ""
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        resume_text += text + "\n"

        elif uploaded_file.name.endswith(".docx"):
            doc = Document(uploaded_file)
            resume_text = ""
            for para in doc.paragraphs:
                resume_text += para.text + "\n"

        else:
            st.error("Unsupported file format")
            st.stop()

        # ===== STEP 1: KEYWORDS =====
        st.info("Extracting keywords...")

        keywords = extract_keyword(jd_text)

        keywords_list = [
            k.strip().lower()
            for k in keywords.split(",")
            if k.strip()
        ]

        keywords_list = list(set(keywords_list))

        st.subheader("Extracted Keywords")
        st.success(", ".join(keywords_list))

        # ===== STEP 2: INITIAL SCORE =====
        resume_lower = resume_text.lower()

        matched = [k for k in keywords_list if k in resume_lower]
        missing = [k for k in keywords_list if k not in resume_lower]

        score = int((len(matched) / len(keywords_list)) * 100) if keywords_list else 0

        st.subheader("Initial ATS Score")
        st.progress(score)
        st.write(f"{score}% match with job description")

        st.subheader("Matched Skills")
        st.success(", ".join(matched) if matched else "None")

        st.subheader("Missing Skills")
        st.warning(", ".join(missing) if missing else "None")

        # ===== STEP 3: OPTIMIZE =====
        st.info("Optimizing resume using AI...")

        optimized_resume = call_local_llm(resume_text, jd_text)

        st.subheader("Optimized Resume")
        st.write(optimized_resume)

        # ===== STEP 4: FINAL SCORE =====
        optimized_lower = optimized_resume.lower()

        matched_new = [k for k in keywords_list if k in optimized_lower]

        score_new = int((len(matched_new) / len(keywords_list)) * 100) if keywords_list else 0

        st.subheader("Improved ATS Score")
        st.progress(score_new)
        st.write(f"{score_new}% match after optimization")

        # ===== STEP 5: DOWNLOAD =====
        st.download_button(
            label="Download Optimized Resume",
            data=optimized_resume,
            file_name="optimized_resume.txt",
            mime="text/plain"
        )